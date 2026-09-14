"""Generate the 3 fictitious sample applications + the FakeLLMProvider answers for each.

Run:  C:\\Users\\User\\anaconda3\\envs\\granter\\python.exe sample_data\\make_samples.py
Every fake quote is copied from the text below; the script re-parses the generated files and checks
that each quote verifies (except the one deliberately fabricated quote in sample 2).
Organizations, people and amounts are invented.
"""
import json
import sys
from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from src import config  # noqa: E402
from src.ingestion import parse_document  # noqa: E402
from src.pipeline import verify_quote  # noqa: E402

# --- 1. Clear case (PDF) ------------------------------------------------------------------------
PDF_PAGES = [
    """Community Housing Alliance
Grant Application to the Riverbend Foundation
Submitted: March 3, 2026

Applicant organization: Community Housing Alliance
Tax status: 501(c)(3) nonprofit organization
Address: 1200 Elm Street, Springfield

1. Request
Community Housing Alliance respectfully requests $250,000 over two years from the Riverbend Foundation. This is a request for project support for our Fair Rent Campaign.

2. Organization
Founded in 2009, Community Housing Alliance works to keep low-income families stably housed. We combine tenant education with policy advocacy at the state capitol.""",
    """3. Project description
The Fair Rent Campaign is a two-year statewide policy advocacy effort to pass tenant protection legislation, including limits on no-cause evictions and a statewide rental assistance fund.

Our primary strategy is policy advocacy: we will draft model legislation, brief legislators, and organize tenant testimony at committee hearings.

The campaign addresses housing stability and economic justice for renters who spend more than half of their income on rent.

4. Population served
The campaign focuses on low-income families who rent their homes, particularly households earning below 50 percent of the area median income.""",
    """5. Geography
The campaign is statewide, with organizers in eight counties and a coordinating office in the state capital.

6. Budget
Total amount requested from the Riverbend Foundation: $250,000 (USD).
Year 1: $130,000. Year 2: $120,000.

7. Expected outcomes
Passage of at least one tenant protection bill, 1,500 tenants trained on their rights, and 300 tenants testifying or meeting with legislators.""",
]

# --- 2. Ambiguous case (DOCX): project vs capacity language, two amounts, implicit population ----
DOCX_TITLE = "Northside Youth Commons"
DOCX_SECTIONS = [
    (None, [
        "Proposal to the Lakeshore Community Fund",
        "Northside Youth Commons is a neighborhood center that has served young people ages 12 to 18 since 2014.",
        "Most of our members attend Hawthorne Middle School and Eastfield High School, where more than 80 percent of students qualify for free or reduced-price lunch.",
    ]),
    ("Why now", [
        "Over the next two years we want to launch Pathways, an after-school digital skills program, while also strengthening our organizational capacity.",
        "Our staff has doubled since 2022, but our finance, evaluation and fundraising systems have not kept pace.",
        "Funding would allow us to pilot Pathways with 120 members and, at the same time, hire a part-time development director and adopt a new data system.",
    ]),
    ("Budget overview", [
        "The total budget for this two-year effort is $1.2 million, funded by public contracts, individual donors and several foundations.",
        "Personnel represents 64 percent of costs; technology and training represent 21 percent.",
    ]),
    ("Our request", [
        "We are requesting $180,000 from the Lakeshore Community Fund over two years.",
        "About two thirds of this grant would support Pathways program delivery and one third would strengthen our internal systems.",
        "Participants will receive mentoring, coding workshops and paid summer internships with local employers.",
    ]),
]
FABRICATED = "The Fund's grant of $180,000 would cover 15 percent of the total budget."

# --- 3. Several plausible strategies (TXT, pages split by form feed) ---------------------------
TXT_PAGES = [
    """GREAT LAKES FAIR FUTURES COALITION
Application for General Operating Support
Submitted to: The Harbor Light Trust
Date: April 14, 2026

About us
Great Lakes Fair Futures Coalition is a network of 34 member organizations in Michigan, Ohio and Wisconsin. We were founded in 2017 by tenant unions, immigrant rights groups, faith congregations and environmental justice organizations.

Our request
We request $400,000 in general operating support over two years. Unrestricted funds let the coalition respond quickly when legislative windows open.
""",
    """How we work
Our work rests on three equal pillars, and no single pillar comes first.
Policy advocacy: we lobby state legislatures and city councils for fair housing, clean air and wage standards.
Community organizing: our organizers recruit and train residents to lead local campaigns in their own neighborhoods.
Coalition building: we convene member organizations to agree on shared platforms across issues and across state lines.

Issues we address
Affordable housing and tenant rights; lead pipe replacement and clean air near highways; fair wages for warehouse and care workers; access to health care for immigrant families; and voting access.
""",
    """Communities we serve
Residents of working-class neighborhoods in Detroit, Cleveland, Toledo and Milwaukee, most of them Black and Latino families, recent immigrants, and hourly workers.

Budget
Our annual budget is $1.9 million. The Harbor Light Trust grant would be 10.5 percent of our two-year operating budget.
""",
]


def q(text, page, kind="explicit"):
    return {"text": text, "page": page, "evidence_type": kind}


def f(value, self_confidence, *quotes):
    return {"value": value, "evidence": list(quotes), "self_confidence": self_confidence}


FAKE = {
    "01_clear_housing_alliance": {
        "organization_name": f("Community Housing Alliance", 0.98,
                               q("Applicant organization: Community Housing Alliance", 1),
                               q("Community Housing Alliance respectfully requests $250,000 over two years", 1)),
        "grant_type": f("Project Support", 0.95,
                        q("This is a request for project support for our Fair Rent Campaign.", 1)),
        "primary_strategy": f("Policy Advocacy", 0.95,
                              q("Our primary strategy is policy advocacy", 2),
                              q("a two-year statewide policy advocacy effort to pass tenant protection legislation", 2)),
        "issues": f(["Housing", "Economic Justice"], 0.92,
                    q("The campaign addresses housing stability and economic justice for renters", 2)),
        "geography": f(["Statewide"], 0.95,
                       q("The campaign is statewide, with organizers in eight counties", 3),
                       q("a two-year statewide policy advocacy effort", 2)),
        "target_population": f(["Low-Income Families"], 0.93,
                               q("The campaign focuses on low-income families who rent their homes", 2),
                               q("works to keep low-income families stably housed", 1)),
        "amount_requested": f(250000, 0.97,
                              q("Community Housing Alliance respectfully requests $250,000 over two years from the Riverbend Foundation.", 1),
                              q("Total amount requested from the Riverbend Foundation: $250,000 (USD).", 3)),
        "project_summary": f(
            "Community Housing Alliance requests $250,000 over two years for the Fair Rent Campaign, a statewide "
            "policy advocacy effort to pass tenant protection legislation, including limits on no-cause evictions "
            "and a statewide rental assistance fund, benefiting low-income renter families.", 0.9,
            q("The Fair Rent Campaign is a two-year statewide policy advocacy effort to pass tenant protection legislation, including limits on no-cause evictions and a statewide rental assistance fund.", 2),
            q("The campaign focuses on low-income families who rent their homes", 2)),
    },
    "02_ambiguous_youth_center": {
        "organization_name": f("Northside Youth Commons", 0.95,
                               q("Northside Youth Commons is a neighborhood center that has served young people ages 12 to 18 since 2014.", 1)),
        "grant_type": f("Project Support", 0.5,
                        q("launch Pathways, an after-school digital skills program, while also strengthening our organizational capacity", 2, "inferred"),
                        q("About two thirds of this grant would support Pathways program delivery", 4, "inferred")),
        "primary_strategy": f("Direct Service", 0.6,
                              q("Participants will receive mentoring, coding workshops and paid summer internships with local employers.", 4, "inferred")),
        "issues": f(["Education", "Workforce Development"], 0.7,
                    q("an after-school digital skills program", 2),
                    q("paid summer internships with local employers", 4, "inferred")),
        "geography": f(["Local"], 0.6,
                       q("Northside Youth Commons is a neighborhood center", 1, "inferred"),
                       q("Most of our members attend Hawthorne Middle School and Eastfield High School", 1, "inferred")),
        "target_population": f(["Children & Youth", "Low-Income Families"], 0.55,
                               q("has served young people ages 12 to 18", 1, "inferred"),
                               q("more than 80 percent of students qualify for free or reduced-price lunch", 1, "inferred")),
        "amount_requested": f(180000, 0.75,
                              q("We are requesting $180,000 from the Lakeshore Community Fund over two years.", 4),
                              q("The total budget for this two-year effort is $1.2 million", 3, "inferred"),
                              q(FABRICATED, 3)),  # not in the document -> "Unverified quote"
        "project_summary": f(
            "Northside Youth Commons requests $180,000 over two years to launch Pathways, an after-school digital "
            "skills program for young people ages 12 to 18, and to strengthen its finance, evaluation and "
            "fundraising systems. The total two-year budget is $1.2 million.", 0.7,
            q("Over the next two years we want to launch Pathways, an after-school digital skills program, while also strengthening our organizational capacity.", 2),
            q("Our staff has doubled since 2022, but our finance, evaluation and fundraising systems have not kept pace.", 2, "inferred")),
    },
    "03_multi_strategy_coalition": {
        "organization_name": f("Great Lakes Fair Futures Coalition", 0.97,
                               q("Great Lakes Fair Futures Coalition is a network of 34 member organizations", 1)),
        "grant_type": f("General Operating Support", 0.95,
                        q("We request $400,000 in general operating support over two years.", 1),
                        q("Application for General Operating Support", 1)),
        "primary_strategy": f("Policy Advocacy", 0.3,
                              q("Our work rests on three equal pillars, and no single pillar comes first.", 2, "inferred"),
                              q("Policy advocacy: we lobby state legislatures and city councils for fair housing, clean air and wage standards.", 2, "inferred"),
                              q("Coalition building: we convene member organizations to agree on shared platforms across issues and across state lines.", 2, "inferred")),
        "issues": f(["Housing", "Environment & Climate", "Economic Justice", "Health", "Civil Rights"], 0.6,
                    q("Affordable housing and tenant rights; lead pipe replacement and clean air near highways; fair wages for warehouse and care workers; access to health care for immigrant families; and voting access.", 2),
                    q("founded in 2017 by tenant unions, immigrant rights groups, faith congregations and environmental justice organizations", 1, "inferred")),
        "geography": f(["Regional (multi-state)"], 0.8,
                       q("a network of 34 member organizations in Michigan, Ohio and Wisconsin", 1, "inferred"),
                       q("across issues and across state lines", 2, "inferred")),
        "target_population": f(["Communities of Color", "Immigrants & Refugees", "Workers", "Low-Income Families"], 0.75,
                               q("Residents of working-class neighborhoods in Detroit, Cleveland, Toledo and Milwaukee, most of them Black and Latino families, recent immigrants, and hourly workers.", 3, "inferred")),
        "amount_requested": f(400000, 0.9,
                              q("We request $400,000 in general operating support over two years.", 1)),
        "project_summary": f(
            "Great Lakes Fair Futures Coalition, a network of 34 member organizations in Michigan, Ohio and "
            "Wisconsin, requests $400,000 in general operating support over two years to sustain policy advocacy, "
            "community organizing and coalition building on housing, environmental, wage, health and voting issues.", 0.8,
            q("Great Lakes Fair Futures Coalition is a network of 34 member organizations in Michigan, Ohio and Wisconsin.", 1),
            q("We request $400,000 in general operating support over two years.", 1),
            q("Our work rests on three equal pillars, and no single pillar comes first.", 2, "inferred")),
    },
}


def write_pdf(path: Path):
    doc = pymupdf.open()
    for text in PDF_PAGES:
        page = doc.new_page()
        assert page.insert_textbox(pymupdf.Rect(72, 72, 540, 760), text, fontsize=11, fontname="helv") >= 0
    doc.save(path)


def write_docx(path: Path):
    d = docx.Document()
    d.add_heading(DOCX_TITLE, level=0)
    for heading, paragraphs in DOCX_SECTIONS:
        if heading:
            d.add_heading(heading, level=1)
        for p in paragraphs:
            d.add_paragraph(p)
    d.save(path)


def main():
    files = {"01_clear_housing_alliance.pdf": write_pdf, "02_ambiguous_youth_center.docx": write_docx,
             "03_multi_strategy_coalition.txt": lambda p: p.write_text("\f".join(TXT_PAGES), encoding="utf-8")}
    out = HERE / "fake_responses"
    out.mkdir(exist_ok=True)
    settings = config.settings()
    for name, writer in files.items():
        path = HERE / name
        writer(path)
        stem = path.stem
        (out / f"{stem}.json").write_text(json.dumps(FAKE[stem], indent=2), encoding="utf-8")
        doc = parse_document(name, path.read_bytes(), settings)
        for field, answer in FAKE[stem].items():
            for quote in answer["evidence"]:
                ok, _ = verify_quote(quote["text"], doc.pages, settings["evidence"]["fuzzy_ratio"])
                assert ok == (quote["text"] != FABRICATED), f"{stem}.{field}: {quote['text']!r}"
        print(f"{name}: {len(doc.pages)} pages/sections, fake answer written")


if __name__ == "__main__":
    main()
