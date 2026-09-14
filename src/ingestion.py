"""PDF/DOCX/TXT bytes -> ParsedDocument. Uploads are untrusted: validate before parsing."""
import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf

OCR_REQUIRED = "OCR required — manual processing needed."
MAGIC = {"pdf": b"%PDF", "docx": b"PK"}


class IngestionError(ValueError):
    pass


@dataclass
class ParsedDocument:
    document_id: str  # sha256 of the raw bytes
    filename: str
    file_type: str    # pdf | docx | txt
    pages: list[str]  # DOCX: sections, not pages


def parse_document(filename: str, data: bytes, settings: dict) -> ParsedDocument:
    cfg = settings["ingestion"]
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in ("pdf", "docx", "txt"):
        raise IngestionError(f"Unsupported file type '.{ext}'. Use PDF, DOCX or TXT.")
    if len(data) > cfg["max_upload_mb"] * 1024 * 1024:
        raise IngestionError(f"File is larger than {cfg['max_upload_mb']} MB.")
    if ext in MAGIC and not data.startswith(MAGIC[ext]):
        raise IngestionError(f"File content does not look like a valid .{ext} file.")

    if ext == "pdf":
        pages = _pdf(data)
        if sum(len(p.strip()) for p in pages) < cfg["min_pdf_text_chars"]:
            # ponytail: sin OCR; conectar aquí un extractor OCR cuando haya PDFs escaneados
            raise IngestionError(OCR_REQUIRED)
    elif ext == "docx":
        pages = _docx(data, cfg["docx_paragraphs_per_section"])
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise IngestionError("TXT file is not valid UTF-8.") from None
        pages = [p for p in text.split("\f") if p.strip()]

    if not pages:
        raise IngestionError("Document contains no text.")
    return ParsedDocument(hashlib.sha256(data).hexdigest(), Path(filename).name, ext, pages)


def _pdf(data: bytes) -> list[str]:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return [page.get_text() for page in doc]
    except Exception as e:  # corrupt / encrypted PDF
        raise IngestionError(f"Could not read PDF ({type(e).__name__}).") from None


def _docx(data: bytes, per_section: int) -> list[str]:
    try:
        paragraphs = docx.Document(io.BytesIO(data)).paragraphs
    except Exception as e:
        raise IngestionError(f"Could not read DOCX ({type(e).__name__}).") from None
    sections, current = [], []
    # ponytail: solo párrafos, las tablas DOCX se ignoran; leer doc.tables si los presupuestos vienen en tablas
    for p in paragraphs:
        if not p.text.strip():
            continue
        is_heading = (p.style.name or "").startswith(("Heading", "Title"))
        if current and (is_heading or len(current) >= per_section):
            sections.append("\n".join(current))
            current = []
        current.append(p.text)
    if current:
        sections.append("\n".join(current))
    return sections
